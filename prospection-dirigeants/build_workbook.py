# -*- coding: utf-8 -*-
"""Construit le classeur de prospection et le fichier de messages LinkedIn.

Usage : python build_workbook.py [dossier_de_sortie]
Sortie : out/Cibles_Mehdi_Kellal_v2.xlsx et out/messages_linkedin.md
"""

import importlib.util
import os
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ICI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ICI)
import messages as msg  # noqa: E402

SEUIL_HAUT, SEUIL_MOYEN = 13, 9
TITRE = PatternFill("solid", fgColor="1F3864")
NOTE = PatternFill("solid", fgColor="FFF2CC")
ENTETE = PatternFill("solid", fgColor="D9E2F3")


def charger_cibles():
    chemin = os.path.join(ICI, "data", "cibles.py")
    spec = importlib.util.spec_from_file_location("cibles", chemin)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return [dict(zip(mod.CHAMPS, c)) for c in mod.CIBLES]


def priorite(total):
    if total >= SEUIL_HAUT:
        return "Haute"
    return "Moyenne" if total >= SEUIL_MOYEN else "Basse"


def largeurs(ws, valeurs):
    for i, l in enumerate(valeurs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = l


def feuille_cibles(wb, cibles):
    ws = wb.create_sheet("Cibles")
    entetes = ["Famille", "Entreprise", "Ce qu'elle fait", "Dépendance publique", "IA / numérique",
               "Santé", "Déploiement territorial", "Stade et levée", "Poste à créer", "Total /18",
               "Priorité", "Signal récent", "Angle pour Mehdi", "Points à vérifier", "Source", "Origine"]
    ws.append(entetes)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = ENTETE
        c.alignment = Alignment(wrap_text=True, vertical="center")

    for cible in cibles:
        n = cible["notes"]
        ligne = ws.max_row + 1
        ws.append([cible["famille"], cible["entreprise"], cible["quoi"], *n, None, None,
                   cible["signal"], cible["angle"], cible["verif"], cible["source"], cible["origine"]])
        ws.cell(ligne, 10).value = f"=SUM(D{ligne}:I{ligne})"
        ws.cell(ligne, 11).value = (
            f'=IF(J{ligne}>=Synthèse!$B$16,"Haute",IF(J{ligne}>=Synthèse!$B$17,"Moyenne","Basse"))')
        for col in range(4, 10):
            ws.cell(ligne, col).fill = NOTE
            ws.cell(ligne, col).alignment = Alignment(horizontal="center")
        ws.cell(ligne, 10).font = Font(bold=True)
        for col in (3, 12, 13, 14):
            ws.cell(ligne, col).alignment = Alignment(wrap_text=True, vertical="top")

    ws.auto_filter.ref = f"A1:P{ws.max_row}"
    ws.freeze_panes = "C2"
    largeurs(ws, [15, 26, 42, 11, 11, 8, 12, 11, 11, 10, 10, 60, 45, 40, 45, 8])
    return ws


def feuille_contacts(wb, cibles):
    ws = wb.create_sheet("Contacts")
    ws.append(["Entreprise", "Priorité", "Nom", "Fonction", "LinkedIn", "Point d'accroche",
               "Statut", "Date de contact", "Réponse", "Suite à donner"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = ENTETE
    for cible in sorted(cibles, key=lambda c: -sum(c["notes"])):
        ws.append([cible["entreprise"], priorite(sum(cible["notes"])), cible["nom"] or "À identifier",
                   cible["fonction"], cible["linkedin"], cible["accroche"], cible["statut"], "", "", ""])
    ws.auto_filter.ref = f"A1:J{ws.max_row}"
    ws.freeze_panes = "A2"
    largeurs(ws, [26, 10, 26, 34, 48, 55, 26, 15, 15, 30])
    for ligne in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        ligne[0].alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def feuille_messages(wb, cibles):
    ws = wb.create_sheet("Messages LinkedIn")
    ws.append(["Entreprise", "Priorité", "Destinataire", "LinkedIn",
               "Note de connexion (300 car. max)", "Caractères", "Premier message (après acceptation)"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = ENTETE
    for cible in sorted(cibles, key=lambda c: -sum(c["notes"])):
        note = msg.note_connexion(cible)
        ws.append([cible["entreprise"], priorite(sum(cible["notes"])),
                   cible["nom"] or "À identifier", cible["linkedin"], note, len(note),
                   msg.message_long(cible)])
    ws.auto_filter.ref = f"A1:G{ws.max_row}"
    ws.freeze_panes = "A2"
    largeurs(ws, [26, 10, 26, 44, 75, 11, 110])
    for ligne in ws.iter_rows(min_row=2):
        for cellule in (ligne[4], ligne[6]):
            cellule.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def feuille_synthese(wb, cibles):
    ws = wb.create_sheet("Synthèse", 0)
    ws["A1"] = "Cibles pour Mehdi Kellal : santé, transformation, IA"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = TITRE
    ws["A2"] = ("Postes visés : directeur des affaires publiques, partenariats publics, "
                "Head of Public Sector. Liste enrichie à 104 cibles en septembre 2026.")
    ws.append([])
    ws.append(["Famille", "Nb cibles", "Priorité haute", "Score moyen /18", "Pitch"])
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = ENTETE

    for famille in ("Santé", "Transformation", "IA"):
        lot = [c for c in cibles if c["famille"] == famille]
        totaux = [sum(c["notes"]) for c in lot]
        ws.append([famille, len(lot), sum(1 for t in totaux if t >= SEUIL_HAUT),
                   round(sum(totaux) / len(totaux), 1), msg.PITCH[famille].capitalize()])
    totaux = [sum(c["notes"]) for c in cibles]
    ws.append(["Total", len(cibles), sum(1 for t in totaux if t >= SEUIL_HAUT),
               round(sum(totaux) / len(totaux), 1), ""])
    ws[f"A{ws.max_row}"].font = Font(bold=True)

    ws.append([])
    ws.append(["Mode d'emploi"])
    ws[f"A{ws.max_row}"].font = Font(bold=True)
    for texte in (
        "• Onglet Cibles : les notes de 0 à 3 (cellules jaunes) sont modifiables ; le total et la priorité se recalculent.",
        "• Priorité : Haute si total ≥ 13, Moyenne si total ≥ 9, sinon Basse (seuils modifiables en B16 et B17).",
        "• Les notes sont une appréciation à partir des informations publiques disponibles ; à ajuster selon les échanges de Mehdi.",
        "• Colonne Origine : v1 = liste initiale, v2 = ajout de septembre 2026.",
        "• Onglet Contacts : dirigeants identifiés via FullEnrich et presse spécialisée, sans enrichissement email ni téléphone.",
        "• Onglet Messages LinkedIn : note de connexion (limite de 300 caractères respectée) et premier message, à relire avant envoi.",
        "• Points à vérifier avant approche : dernière levée, présence d'un profil affaires publiques en interne, dirigeant actuel.",
    ):
        ws.append([texte])
    ws.append([])
    ws["A16"] = "Seuil priorité haute"
    ws["B16"] = SEUIL_HAUT
    ws["A17"] = "Seuil priorité moyenne"
    ws["B17"] = SEUIL_MOYEN
    ws["A19"] = f"Classeur généré le {date.today().strftime('%d/%m/%Y')} par prospection-dirigeants/build_workbook.py"
    largeurs(ws, [46, 14, 16, 18, 80])
    return ws


def feuille_sources(wb, cibles):
    ws = wb.create_sheet("Sources")
    ws.append(["Entreprise", "Sujet", "Lien"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = ENTETE
    for cible in cibles:
        if cible["source"]:
            ws.append([cible["entreprise"], msg.accroche(cible["signal"], 160), cible["source"]])
    ws.append(["Toutes", "Dirigeants identifiés par recherche FullEnrich (16/09/2026)", ""])
    ws.append(["Toutes", "Recherches web de septembre 2026 (presse spécialisée, communiqués)", ""])
    largeurs(ws, [26, 90, 90])
    return ws


def ecrire_messages_md(cibles, dossier):
    lignes = ["# Messages LinkedIn — prospection dirigeants",
              "",
              f"Généré le {date.today().strftime('%d/%m/%Y')}. "
              "Relire et adapter avant envoi : ces messages sont des bases, pas des envois automatiques.",
              ""]
    for cible in sorted(cibles, key=lambda c: (-sum(c["notes"]), c["entreprise"])):
        total = sum(cible["notes"])
        note = msg.note_connexion(cible)
        lignes += [f"## {cible['entreprise']} — {cible['nom'] or 'dirigeant à identifier'}"
                   f" ({cible['fonction']})",
                   "",
                   f"*{cible['famille']} · priorité {priorite(total).lower()} · {total}/18*"
                   + (f" · [profil]({cible['linkedin']})" if cible["linkedin"] else ""),
                   "",
                   f"**Note de connexion ({len(note)} caractères)**",
                   "",
                   f"> {note}",
                   "",
                   "**Premier message**",
                   "",
                   "```",
                   msg.message_long(cible),
                   "```",
                   ""]
    chemin = os.path.join(dossier, "messages_linkedin.md")
    with open(chemin, "w", encoding="utf-8") as f:
        f.write("\n".join(lignes))
    return chemin


def main():
    dossier = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ICI, "out")
    os.makedirs(dossier, exist_ok=True)
    cibles = charger_cibles()

    wb = Workbook()
    wb.remove(wb.active)
    feuille_cibles(wb, cibles)
    feuille_contacts(wb, cibles)
    feuille_messages(wb, cibles)
    feuille_sources(wb, cibles)
    feuille_synthese(wb, cibles)
    wb._sheets = [wb["Synthèse"], wb["Cibles"], wb["Contacts"], wb["Messages LinkedIn"], wb["Sources"]]

    xlsx = os.path.join(dossier, "Cibles_Mehdi_Kellal_v2.xlsx")
    wb.save(xlsx)
    md = ecrire_messages_md(cibles, dossier)

    hautes = sum(1 for c in cibles if sum(c["notes"]) >= SEUIL_HAUT)
    a_completer = sum(1 for c in cibles if not c["nom"])
    print(f"{len(cibles)} cibles — {hautes} en priorité haute — {a_completer} dirigeants à identifier")
    print(xlsx)
    print(md)


if __name__ == "__main__":
    main()
