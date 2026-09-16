# -*- coding: utf-8 -*-
"""Sélectionne les prochaines cibles à contacter.

Relit le classeur (onglets Contacts et Messages LinkedIn) et sort la liste des
dirigeants à approcher cette semaine : priorité haute d'abord, en écartant
ceux qui ont déjà une date de contact et ceux dont le dirigeant reste à
identifier. Écrit out/a_contacter.md.

Usage : python selection.py [classeur.xlsx] [nombre]
"""

import os
import sys
from datetime import date

from openpyxl import load_workbook

ICI = os.path.dirname(os.path.abspath(__file__))
ORDRE = {"Haute": 0, "Moyenne": 1, "Basse": 2}


def lire(classeur):
    wb = load_workbook(classeur)
    contacts = {}
    for rang, ligne in enumerate(wb["Contacts"].iter_rows(min_row=2, values_only=True)):
        entreprise, prio, nom, fonction, linkedin, accroche, statut, contacte = ligne[:8]
        contacts[entreprise] = dict(rang=rang, entreprise=entreprise, priorite=prio, nom=nom,
                                    fonction=fonction, linkedin=linkedin, accroche=accroche,
                                    statut=statut, contacte=contacte)
    for ligne in wb["Messages LinkedIn"].iter_rows(min_row=2, values_only=True):
        entreprise, _, _, _, note = ligne[:5]
        if entreprise in contacts:
            contacts[entreprise]["note"] = note
    return list(contacts.values())


def a_contacter(contacts, combien):
    restants = [c for c in contacts
                if not c["contacte"] and c["nom"] and c["nom"] != "À identifier"]
    # le classeur est déjà trié par score décroissant : on conserve cet ordre
    restants.sort(key=lambda c: (ORDRE.get(c["priorite"], 3), c["rang"]))
    return restants[:combien]


def main():
    classeur = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ICI, "out", "Cibles_Mehdi_Kellal_v2.xlsx")
    combien = int(sys.argv[2]) if len(sys.argv) > 2 else 10
    contacts = lire(classeur)
    lot = a_contacter(contacts, combien)

    lignes = [f"# À contacter — semaine du {date.today().strftime('%d/%m/%Y')}", "",
              f"{len(lot)} dirigeants sélectionnés sur {len(contacts)} cibles "
              f"(priorité haute d'abord, hors contacts déjà engagés).", ""]
    for c in lot:
        lignes += [f"## {c['nom']} — {c['fonction']}, {c['entreprise']} ({str(c['priorite']).lower()})",
                   "",
                   f"- Profil : {c['linkedin'] or 'à retrouver sur LinkedIn'}",
                   f"- Accroche : {c['accroche'] or '—'}",
                   "",
                   f"> {c.get('note', '(note de connexion à générer)')}",
                   ""]
    manquants = [c["entreprise"] for c in contacts if not c["nom"] or c["nom"] == "À identifier"]
    if manquants:
        lignes += ["## Dirigeants encore à identifier", "",
                   ", ".join(sorted(manquants)), ""]

    chemin = os.path.join(os.path.dirname(classeur), "a_contacter.md")
    with open(chemin, "w", encoding="utf-8") as f:
        f.write("\n".join(lignes))
    print(f"{len(lot)} cibles sélectionnées — {len(manquants)} dirigeants à identifier")
    print(chemin)


if __name__ == "__main__":
    main()
